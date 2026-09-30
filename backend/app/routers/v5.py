from __future__ import annotations
from typing import Any
from datetime import datetime, timezone
from io import BytesIO
import json, os, pathlib, uuid
from fastapi import APIRouter, Depends, HTTPException, Body, UploadFile, File, Form, Query
from fastapi.responses import StreamingResponse, HTMLResponse, FileResponse
from sqlalchemy import text
from sqlalchemy.orm import Session
from openpyxl import Workbook, load_workbook
from ..db import get_db
from ..deps import current_user, require_permission, has_permission
from ..security import hash_password, verify_password

router=APIRouter(prefix='/api/v5',tags=['Ambulans Yönetim Sistemi AYS V14'])

MODULES=['Analiz','Araçlar','Arıza & İş Emri','Teknik Destek Talepleri','Uyarılar','Bakım & Servis','Trafik Kazaları','Trafik Cezaları','Tıbbi Donanım','Teknik Birim Deposu','Ambulans Değişimi','Devir/Terkin/Dönüşüm','İstasyonlar & Birimler','Servisler','Rapor Merkezi','Kullanıcılar','Toplu Veri Aktarımı','Ayarlar']
ACTIONS=['GORUNTULE','EKLE','DUZENLE','PASIFE_AL','ONAYLA','RAPORLA']

def rows(r): return [dict(x._mapping) for x in r]

def audit(db,u,etype,eid,islem,yeni=None,gerekce=None):
    db.execute(text("INSERT INTO aftys.audit_log(kullanici_id,entity_type,entity_id,islem,yeni_json,gerekce) VALUES(:u,:t,:id,:i,CAST(:j AS jsonb),:g)"),{'u':str(u.kullanici_id),'t':etype,'id':str(eid),'i':islem,'j':json.dumps(yeni or {},ensure_ascii=False),'g':gerekce})

def next_number(db:Session,series:str,year:int|None=None)->str:
    year=year or datetime.now(timezone.utc).year
    row=db.execute(text("""INSERT INTO aftys.numara_sayaclari(seri,yil,son_no) VALUES(:s,:y,1)
      ON CONFLICT(seri,yil) DO UPDATE SET son_no=aftys.numara_sayaclari.son_no+1
      RETURNING son_no"""),{'s':series,'y':year}).scalar_one()
    return f"{series}-{year}-{int(row):04d}"

def file_store(upload:UploadFile, max_mb:int=25):
    root=pathlib.Path(os.environ.get('AFTYS_UPLOAD_DIR','/data/uploads')); root.mkdir(parents=True,exist_ok=True)
    suffix=pathlib.Path(upload.filename or '').suffix[:12]; key=f"{uuid.uuid4()}{suffix}"; target=root/key
    return root,key,target

@router.get('/me')
def me(db:Session=Depends(get_db),u=Depends(current_user)):
    roles=[r[0] for r in db.execute(text("SELECT r.ad FROM aftys.kullanici_rolleri kr JOIN aftys.roller r ON r.rol_id=kr.rol_id WHERE kr.kullanici_id=:u ORDER BY r.ad"),{'u':str(u.kullanici_id)}).all()]
    perms={}
    for m in MODULES:
        perms[m]={a:has_permission(db,u,m,a) for a in ACTIONS}
    pref=db.execute(text("SELECT bildirimler,tema FROM aftys.kullanici_tercihleri WHERE kullanici_id=:u"),{'u':str(u.kullanici_id)}).mappings().first()
    force=db.execute(text("SELECT COALESCE(parola_degistir_zorunlu,false) FROM aftys.kullanicilar WHERE kullanici_id=:u"),{'u':str(u.kullanici_id)}).scalar()
    return {'kullanici_id':str(u.kullanici_id),'kullanici_adi':u.kullanici_adi,'tc_kimlik_no':db.execute(text("SELECT tc_kimlik_no FROM aftys.kullanicilar WHERE kullanici_id=:u"),{'u':str(u.kullanici_id)}).scalar(),'ad_soyad':u.ad_soyad,'eposta':getattr(u,'eposta',None),'telefon':getattr(u,'telefon',None),'roller':roles,'yetkiler':perms,'parola_degistir_zorunlu':bool(force),'tercihler':dict(pref) if pref else {'bildirimler':{},'tema':'SISTEM'}}

@router.put('/me/profile')
def update_profile(b:dict[str,Any]=Body(...),db:Session=Depends(get_db),u=Depends(current_user)):
    db.execute(text("UPDATE aftys.kullanicilar SET ad_soyad=:a,eposta=:e,telefon=:t,updated_at=now() WHERE kullanici_id=:u"),{'a':str(b.get('ad_soyad') or u.ad_soyad).strip(),'e':str(b.get('eposta') or '').strip() or None,'t':str(b.get('telefon') or '').strip() or None,'u':str(u.kullanici_id)})
    audit(db,u,'KULLANICI',u.kullanici_id,'PROFILE_UPDATE',{'ad_soyad':b.get('ad_soyad'),'eposta':b.get('eposta'),'telefon':b.get('telefon')}); db.commit(); return {'ok':True}

@router.put('/me/password')
def change_password(b:dict[str,Any]=Body(...),db:Session=Depends(get_db),u=Depends(current_user)):
    old=str(b.get('eski_parola') or ''); new=str(b.get('yeni_parola') or '')
    force=bool(db.execute(text("SELECT COALESCE(parola_degistir_zorunlu,false) FROM aftys.kullanicilar WHERE kullanici_id=:u"),{'u':str(u.kullanici_id)}).scalar())
    if not force and not verify_password(old,u.parola_hash): raise HTTPException(422,'Mevcut parola hatalı')
    if len(new)<8: raise HTTPException(422,'Yeni parola en az 8 karakter olmalı')
    db.execute(text("UPDATE aftys.kullanicilar SET parola_hash=:p,parola_degistir_zorunlu=false,updated_at=now() WHERE kullanici_id=:u"),{'p':hash_password(new),'u':str(u.kullanici_id)}); audit(db,u,'KULLANICI',u.kullanici_id,'PASSWORD_CHANGE'); db.commit(); return {'ok':True}

@router.put('/me/preferences')
def preferences(b:dict[str,Any]=Body(...),db:Session=Depends(get_db),u=Depends(current_user)):
    bild=b.get('bildirimler') or {}; tema=b.get('tema') or 'SISTEM'
    db.execute(text("""INSERT INTO aftys.kullanici_tercihleri(kullanici_id,bildirimler,tema) VALUES(:u,CAST(:b AS jsonb),:t)
      ON CONFLICT(kullanici_id) DO UPDATE SET bildirimler=EXCLUDED.bildirimler,tema=EXCLUDED.tema,updated_at=now()"""),{'u':str(u.kullanici_id),'b':json.dumps(bild,ensure_ascii=False),'t':tema}); db.commit(); return {'ok':True}

@router.post('/feedback',status_code=201)
def feedback(b:dict[str,Any]=Body(...),db:Session=Depends(get_db),u=Depends(current_user)):
    tur=str(b.get('tur') or 'ONERI'); konu=str(b.get('konu') or '').strip(); mesaj=str(b.get('mesaj') or '').strip()
    if not konu or not mesaj: raise HTTPException(422,'Konu ve açıklama zorunlu')
    rid=db.execute(text("INSERT INTO aftys.geri_bildirimler(kullanici_id,tur,konu,mesaj) VALUES(:u,:t,:k,:m) RETURNING geri_bildirim_id"),{'u':str(u.kullanici_id),'t':tur,'k':konu,'m':mesaj}).scalar_one(); db.commit(); return {'geri_bildirim_id':rid}

@router.put('/admin/users/{uid}/profile')
def admin_edit_user(uid:uuid.UUID,b:dict[str,Any]=Body(...),db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Kullanıcılar','DUZENLE')
    if not db.execute(text("SELECT 1 FROM aftys.kullanicilar WHERE kullanici_id=:id"),{'id':str(uid)}).first(): raise HTTPException(404,'Kullanıcı bulunamadı')
    db.execute(text("UPDATE aftys.kullanicilar SET ad_soyad=:a,eposta=:e,telefon=:t,aktif_mi=:ak,updated_at=now() WHERE kullanici_id=:id"),{'a':str(b.get('ad_soyad') or '').strip(),'e':str(b.get('eposta') or '').strip() or None,'t':str(b.get('telefon') or '').strip() or None,'ak':bool(b.get('aktif_mi',True)),'id':str(uid)})
    audit(db,u,'KULLANICI',uid,'PROFILE_ADMIN_UPDATE',b); db.commit(); return {'ok':True}

@router.get('/talepler')
def list_requests(db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Teknik Destek Talepleri','GORUNTULE')
    return rows(db.execute(text("""SELECT t.*,trim(concat(COALESCE(ib.kod,''),CASE WHEN ib.kod IS NOT NULL AND ib.telsiz_adi IS NOT NULL THEN ' ' ELSE '' END,COALESCE(ib.telsiz_adi,''),' — ',ib.ad)) istasyon_birim,(SELECT deger FROM aftys.arac_kimlik_gecmisi WHERE arac_id=t.arac_id AND alan_tipi='PLAKA' AND bitis_tarihi IS NULL LIMIT 1) plaka,
      k.ad_soyad kaydeden FROM aftys.teknik_talepler t LEFT JOIN aftys.istasyon_birimler ib ON ib.birim_id=t.istasyon_birim_id LEFT JOIN aftys.kullanicilar k ON k.kullanici_id=t.created_by WHERE COALESCE(t.pasif_mi,false)=false ORDER BY t.created_at DESC LIMIT 500""")))

@router.post('/talepler',status_code=201)
def create_request(b:dict[str,Any]=Body(...),db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Teknik Destek Talepleri','EKLE')
    if not b.get('istasyon_birim_id') or not b.get('talep_turu') or not str(b.get('aciklama') or '').strip(): raise HTTPException(422,'İstasyon/Birim, talep türü ve açıklama zorunlu')
    is_fault=str(b.get('talep_turu') or '').upper()=='ARIZA'
    fault_type=str(b.get('ariza_kategori') or '') if is_fault else None
    if is_fault and fault_type not in ('ARAC_DONANIM','TIBBI_DONANIM'): raise HTTPException(422,'Arıza türü seçilmelidir')
    solution=str(b.get('ariza_cozum_yeri') or '') if is_fault else None
    if is_fault and solution not in ('SERVIS','TEKNIK_DESTEK'): raise HTTPException(422,'Arıza çözüm yeri seçilmelidir')
    no=next_number(db,'TLP')
    sql="""INSERT INTO aftys.teknik_talepler(talep_no,istasyon_birim_id,bildiren_kisi,iletisim_yontemi,talep_turu,arac_id,aciklama,ariza_kategori,ariza_sistem_parca,arac_gorev_yapabiliyor,aciliyet,durum,created_by,bildiren_kullanici_id,ariza_cozum_yeri,servis_id,gayrifaal_baslangic) VALUES(:n,:b,:bk,:iy,:tt,:a,:ac,:kat,NULL,false,'NORMAL',:d,:u,:bu,:cy,:sid,CASE WHEN :fault THEN now() ELSE NULL END) RETURNING talep_id"""
    rid=db.execute(text(sql),{'n':no,'b':b['istasyon_birim_id'],'bk':b.get('bildiren_kisi') or None,'iy':b.get('iletisim_yontemi') or 'TELEFON','tt':b['talep_turu'],'a':b.get('arac_id') or None,'ac':b['aciklama'],'kat':fault_type,'d':'GAYRIFAAL' if is_fault else 'YENI','u':str(u.kullanici_id),'bu':b.get('bildiren_kullanici_id') or None,'cy':solution,'sid':b.get('servis_id') or None,'fault':is_fault}).scalar_one()
    if is_fault and b.get('arac_id'):
        did=db.execute(text("SELECT arac_durum_id FROM aftys.arac_durumlari WHERE kod='ARIZALI_SERVISTE'")).scalar_one_or_none()
        if did: db.execute(text("UPDATE aftys.araclar SET arac_durum_id=:d WHERE arac_id=:a"),{'d':did,'a':b['arac_id']})
    initial='GAYRIFAAL' if is_fault else 'YENI'; note='Arıza bildirimi alındı; gayrifaal süreci otomatik başlatıldı' if is_fault else 'Talep / çağrı kaydı oluşturuldu'
    db.execute(text("INSERT INTO aftys.teknik_talep_zaman_cizelgesi(talep_id,durum,aciklama,kullanici_id) VALUES(:id,:d,:n,:u)"),{'id':rid,'d':initial,'n':note,'u':str(u.kullanici_id)})
    audit(db,u,'TEKNIK_TALEP',rid,'CREATE',{'talep_no':no,'talep_turu':b['talep_turu'],'ariza_turu':fault_type,'ariza_cozum_yeri':solution}); db.commit(); return {'talep_id':rid,'talep_no':no}

@router.put('/talepler/{talep_id}')
def update_request(talep_id:int,b:dict[str,Any]=Body(...),db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Teknik Destek Talepleri','DUZENLE')
    current=db.execute(text("SELECT * FROM aftys.teknik_talepler WHERE talep_id=:id"),{'id':talep_id}).mappings().first()
    if not current: raise HTTPException(404,'Talep bulunamadı')
    fields=['istasyon_birim_id','bildiren_kisi','iletisim_yontemi','talep_turu','arac_id','aciklama','ariza_kategori','ariza_cozum_yeri','servis_id']
    vals={k:(b[k] if k in b else current.get(k)) for k in fields}; vals['id']=talep_id
    if not vals['istasyon_birim_id'] or not vals['talep_turu'] or not str(vals['aciklama'] or '').strip(): raise HTTPException(422,'İstasyon/Birim, talep türü ve açıklama zorunlu')
    db.execute(text("""UPDATE aftys.teknik_talepler SET istasyon_birim_id=:istasyon_birim_id,bildiren_kisi=:bildiren_kisi,iletisim_yontemi=:iletisim_yontemi,talep_turu=:talep_turu,arac_id=:arac_id,aciklama=:aciklama,ariza_kategori=:ariza_kategori,ariza_cozum_yeri=:ariza_cozum_yeri,servis_id=:servis_id,updated_at=now() WHERE talep_id=:id"""),vals)
    audit(db,u,'TEKNIK_TALEP',talep_id,'UPDATE',{'degisiklikler':b}); db.commit(); return {'ok':True}

@router.post('/talepler/{talep_id}/is-emri',status_code=201)
def request_to_workorder(talep_id:int,db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Arıza & İş Emri','EKLE')
    t=db.execute(text("SELECT * FROM aftys.teknik_talepler WHERE talep_id=:id"),{'id':talep_id}).mappings().first()
    if not t: raise HTTPException(404,'Talep bulunamadı')
    if not t['arac_id']: raise HTTPException(422,'İş emri için talepte araç seçilmelidir')
    existing=db.execute(text("SELECT is_emri_id,is_emri_no FROM aftys.is_emirleri WHERE kaynak_tipi='TEKNIK_TALEP' AND kaynak_id=:id"),{'id':talep_id}).mappings().first()
    if existing: return dict(existing)
    no=next_number(db,'IE')
    wid=db.execute(text("""INSERT INTO aftys.is_emirleri(is_emri_no,arac_id,kaynak_tipi,kaynak_id,islem_turu,talep,oncelik,sorumlu_kullanici_id,created_by)
      VALUES(:n,:a,'TEKNIK_TALEP',:kid,:tur,:tal,:on,:u,:u) RETURNING is_emri_id"""),{'n':no,'a':str(t['arac_id']),'kid':talep_id,'tur':'ARIZA' if t['talep_turu']=='ARIZA' else 'DIGER','tal':t['aciklama'],'on':t['aciliyet'] or 'NORMAL','u':str(u.kullanici_id)}).scalar_one()
    if t['talep_turu']=='ARIZA':
        arid=db.execute(text("INSERT INTO aftys.arizalar(arac_id,kategori,sistem_parca,aciklama,durum) VALUES(:a,:k,:s,:ac,'IS_EMRI_ACILDI') RETURNING ariza_id"),{'a':str(t['arac_id']),'k':t['ariza_kategori'] or 'ARAC_DONANIM','s':t['ariza_sistem_parca'],'ac':t['aciklama']}).scalar_one()
        db.execute(text("UPDATE aftys.is_emirleri SET kaynak_tipi='ARIZA',kaynak_id=:ar WHERE is_emri_id=:w"),{'ar':arid,'w':wid})
    db.execute(text("UPDATE aftys.teknik_talepler SET durum='IS_EMRINE_DONUSTURULDU',is_emri_id=:w,updated_at=now() WHERE talep_id=:id"),{'w':wid,'id':talep_id})
    audit(db,u,'IS_EMRI',wid,'CREATE_FROM_TALEP',{'is_emri_no':no,'talep_id':talep_id}); db.commit(); return {'is_emri_id':wid,'is_emri_no':no}

@router.get('/is-emirleri/{is_emri_id}/print',response_class=HTMLResponse)
def print_workorder(is_emri_id:int,db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Arıza & İş Emri','GORUNTULE')
    r=db.execute(text("""SELECT ie.*,a.marka,a.model,a.model_yili,a.gosterge_km,(SELECT deger FROM aftys.arac_kimlik_gecmisi WHERE arac_id=a.arac_id AND alan_tipi='PLAKA' AND bitis_tarihi IS NULL LIMIT 1) plaka,
      ib.ad istasyon,ib.kod telsiz_kodu,s.ad servis_adi,k.ad_soyad teknik_personel
      FROM aftys.is_emirleri ie JOIN aftys.araclar a ON a.arac_id=ie.arac_id LEFT JOIN aftys.istasyon_birimler ib ON ib.birim_id=a.gunluk_birim_id
      LEFT JOIN aftys.servis_islemleri si ON si.kaynak_tipi='IS_EMRI' AND si.kaynak_id=ie.is_emri_id AND si.cikis_tarihi IS NULL LEFT JOIN aftys.servisler s ON s.servis_id=si.servis_id
      LEFT JOIN aftys.kullanicilar k ON k.kullanici_id=ie.sorumlu_kullanici_id WHERE ie.is_emri_id=:id"""),{'id':is_emri_id}).mappings().first()
    if not r: raise HTTPException(404,'İş emri bulunamadı')
    def esc(v): return str(v or '').replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')
    dt=r['acilis_tarihi'].astimezone().strftime('%d.%m.%Y %H:%M') if r['acilis_tarihi'] else ''
    html=f'''<!doctype html><html><head><meta charset="utf-8"><title>{esc(r['is_emri_no'])}</title><style>@page{{size:A5 landscape;margin:5mm}}body{{font:10px Arial;margin:0}}table{{width:100%;border-collapse:collapse}}td,th{{border:1px solid #111;padding:3px}}.c{{text-align:center}}.b{{font-weight:700}}.h{{height:54px}}.sig{{height:55px;vertical-align:top}}@media print{{button{{display:none}}}}</style></head><body><button onclick="print()">Yazdır</button><table><tr><td rowspan="2" class="c b" style="width:15%">T.C.<br>SAĞLIK BAKANLIĞI</td><td class="c b">T.C. SAĞLIK BAKANLIĞI<br>İZMİR AMBULANS SERVİSİ BAŞHEKİMLİĞİ</td><td rowspan="2" class="c b" style="width:15%">AMBULANS<br>YÖNETİMİ</td></tr><tr><td class="c b">TEKNİK DESTEK BİRİMİ<br>ARAÇ ARIZA TESPİT VE ONARIM FORMU</td></tr><tr><td>Doküman Kodu<br><b>AS.FR.07</b></td><td class="c">Yayın Tarihi: 12.02.2018 &nbsp; Revizyon No: 06 &nbsp; Revizyon Tarihi: 05.02.2025</td><td class="c">Sayfa 1/1</td></tr></table><table><tr><td>Teknik Bölüm Protokol / İş Emri No</td><td class="b">{esc(r['is_emri_no'])}</td></tr><tr><td>Teknik Bölüme Bildirim Tarih ve Saati</td><td>{dt}</td></tr><tr><td>Plaka</td><td>{esc(r['plaka'])}</td></tr><tr><td>Görev Yaptığı İstasyon Adı ve Telsiz Kodu</td><td>{esc(r['istasyon'])} {esc(r['telsiz_kodu'])}</td></tr><tr><td>Marka ve Model Yılı</td><td>{esc(r['marka'])} {esc(r['model'])} / {esc(r['model_yili'])}</td></tr><tr><td>Onarım Giriş Km</td><td>{esc(r['gosterge_km'])}</td></tr><tr><td>Servis Adı</td><td>{esc(r['servis_adi'])}</td></tr></table><table><tr><th>TAHMİN EDİLEN ARIZALAR (İstasyon/Teknik Destek Birimi)</th><th>ARIZA TESPİTİ (Servis)</th></tr><tr><td class="h">{esc(r['talep'])}</td><td class="h">{esc(r['yapilan_islem'])}</td></tr><tr><th>İSTASYON PERSONELİ</th><th>TEKNİK DESTEK PERSONELİ / SERVİS YETKİLİSİ</th></tr><tr><td class="sig"></td><td class="sig">{esc(r['teknik_personel'])}</td></tr></table></body></html>'''
    return HTMLResponse(html)

@router.post('/ambulans-degisimi/{degisim_id}/foto',status_code=201)
async def exchange_photo(degisim_id:int,taraf:str=Form(...),plaka:str=Form(''),dosya:UploadFile=File(...),db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Ambulans Değişimi','DUZENLE')
    if taraf not in ('CIKAN','GELEN'): raise HTTPException(422,'Taraf CIKAN veya GELEN olmalı')
    if not db.execute(text("SELECT 1 FROM aftys.ambulans_degisimleri WHERE degisim_id=:id"),{'id':degisim_id}).first(): raise HTTPException(404,'Değişim bulunamadı')
    root,key,target=file_store(dosya); total=0
    with target.open('wb') as f:
        while True:
            chunk=await dosya.read(1024*1024)
            if not chunk: break
            total+=len(chunk)
            if total>25*1024*1024: f.close(); target.unlink(missing_ok=True); raise HTTPException(413,'Dosya 25 MB sınırını aşıyor')
            f.write(chunk)
    eid=db.execute(text("INSERT INTO aftys.dosyalar_ekler(entity_type,entity_id,dosya_adi,mime_type,storage_key,aciklama,yukleyen_id) VALUES(:t,:id,:ad,:m,:k,:ac,:u) RETURNING ek_id"),{'t':f'AMBULANS_DEGISIM_{taraf}','id':str(degisim_id),'ad':dosya.filename or key,'m':dosya.content_type or 'application/octet-stream','k':key,'ac':f'Plaka: {plaka}','u':str(u.kullanici_id)}).scalar_one(); db.commit(); return {'ek_id':str(eid)}

@router.get('/ambulans-degisimi/{degisim_id}/foto')
def exchange_photos(degisim_id:int,db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Ambulans Değişimi','GORUNTULE')
    return rows(db.execute(text("SELECT ek_id::text ek_id,entity_type,dosya_adi,aciklama,yukleme_tarihi FROM aftys.dosyalar_ekler WHERE entity_id=:id AND entity_type IN ('AMBULANS_DEGISIM_CIKAN','AMBULANS_DEGISIM_GELEN') ORDER BY yukleme_tarihi"),{'id':str(degisim_id)}))

TEMPLATES={
 'araclar':['plaka','arac_sinifi','marka','model','model_yili','yakit_turu','baslangic_km','sasi_no','motor_no','adblue_var','tibbi_donanim_markasi'],
 'istasyonlar':['istasyon_adi','istasyon_telsiz_kodu','istasyon_telsiz_adi'],
 'servisler':['servis_turu','ad','hizmet_verilen_markalar','eposta','adres'],
 'kullanicilar':['tc_kimlik_no','ad','soyad','eposta','telefon'],
 'cagri_kayit':['cagri_tarihi_saati','istasyon_telsiz_kodu','plaka','talep_turu','ariza_turu','ariza_cozum_yeri','bildiren_kisi','aciklama','durum']
}

@router.get('/import/{entity}/template')
def import_template(entity:str,db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Toplu Veri Aktarımı','GORUNTULE')
    headers=TEMPLATES.get(entity)
    if not headers: raise HTTPException(404,'Şablon bulunamadı')
    wb=Workbook(); ws=wb.active; ws.title='Veriler'; ws.append(headers)
    examples={'araclar':['35ABC112','AMBULANS','FORD','TRANSIT',2026,'Dizel',0,'SASI123','MOTOR123','EVET',''], 'istasyonlar':['1 No.lu ASHİ','KOD-1','MERKEZ-1'], 'servisler':['ARAC','Örnek Servis','Ford','',''], 'kullanicilar':['12345678901','Örnek','Personel','',''], 'cagri_kayit':['23.09.2026 09:05','555','35ABC112','ARIZA','ARAC_DONANIM','SERVIS','Örnek Personel','Örnek geçmiş arıza','TAMAMLANDI']}
    ws.append(examples[entity]); ws.freeze_panes='A2'
    bio=BytesIO(); wb.save(bio); bio.seek(0)
    return StreamingResponse(bio,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="AmbulansYonetimi_{entity}_sablon.xlsx"'})

def parse_xlsx(upload:UploadFile):
    data=upload.file.read(); wb=load_workbook(BytesIO(data),data_only=True); ws=wb.active
    headers=[str(x.value or '').strip() for x in ws[1]]; out=[]
    for idx,row in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
        if not any(v not in (None,'') for v in row): continue
        out.append((idx,{headers[i]:(row[i] if i<len(row) else None) for i in range(len(headers))}))
    return headers,out

@router.post('/import/{entity}')
def bulk_import(entity:str,dry_run:bool=Query(True),dosya:UploadFile=File(...),db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Toplu Veri Aktarımı','EKLE')
    expected=TEMPLATES.get(entity)
    if not expected: raise HTTPException(404,'Aktarım türü bulunamadı')
    headers,data=parse_xlsx(dosya); missing=[x for x in expected if x not in headers]
    if missing: raise HTTPException(422,'Şablon alanları eksik: '+', '.join(missing))
    errors=[]; valid=[]
    for line,r in data:
        try:
            if entity=='araclar':
                for f in ('plaka','arac_sinifi','marka'):
                    if not str(r.get(f) or '').strip(): raise ValueError(f'{f} zorunlu')
                if db.execute(text("SELECT 1 FROM aftys.arac_kimlik_gecmisi WHERE alan_tipi='PLAKA' AND upper(deger)=upper(:p) AND bitis_tarihi IS NULL"),{'p':str(r['plaka']).strip()}).first(): raise ValueError('Plaka zaten kayıtlı')
            elif entity=='istasyonlar':
                if not str(r.get('istasyon_adi') or '').strip(): raise ValueError('İstasyon adı zorunlu')
                if r.get('istasyon_telsiz_kodu') and db.execute(text("SELECT 1 FROM aftys.istasyon_birimler WHERE kod=:k"),{'k':str(r['istasyon_telsiz_kodu']).strip()}).first(): raise ValueError('İstasyon telsiz kodu zaten kayıtlı')
            elif entity=='servisler':
                if not str(r.get('ad') or '').strip(): raise ValueError('Servis adı zorunlu')
                raw=str(r.get('servis_turu') or 'ARAC').strip().upper().replace('İ','I').replace('Ç','C').replace(' ','_')
                aliases={'ARAC_SERVISI':'ARAC','ARAC':'ARAC','TIBBI_DONANIM_SERVISI':'TIBBI_DONANIM','TIBBI_DONANIM':'TIBBI_DONANIM'}
                if raw not in aliases: raise ValueError('Servis türü ARAC veya TIBBI_DONANIM olmalı')
                r['_servis_turu']=aliases[raw]
                if db.execute(text("SELECT 1 FROM aftys.servisler WHERE lower(ad)=lower(:a) AND aktif_mi=true"),{'a':str(r['ad']).strip()}).first(): raise ValueError('Servis adı zaten kayıtlı')
            elif entity=='cagri_kayit':
                if not str(r.get('istasyon_telsiz_kodu') or '').strip(): raise ValueError('İstasyon telsiz kodu zorunlu')
                if not str(r.get('aciklama') or '').strip(): raise ValueError('Açıklama zorunlu')
                tt=str(r.get('talep_turu') or '').strip().upper()
                if tt=='ARIZA' and str(r.get('ariza_turu') or '').strip().upper() not in ('ARAC_DONANIM','TIBBI_DONANIM'): raise ValueError('Arıza türü geçersiz')
            elif entity=='kullanicilar':
                tc=str(r.get('tc_kimlik_no') or '').strip()
                if len(tc)!=11 or not tc.isdigit(): raise ValueError('TC Kimlik No 11 haneli ve sadece rakam olmalı')
                if not str(r.get('ad') or '').strip() or not str(r.get('soyad') or '').strip(): raise ValueError('ad ve soyad zorunlu')
                if db.execute(text("SELECT 1 FROM aftys.kullanicilar WHERE tc_kimlik_no=:k"),{'k':tc}).first(): raise ValueError('TC Kimlik No zaten kayıtlı')
            valid.append((line,r))
        except Exception as e: errors.append({'satir':line,'hata':str(e)})
    if dry_run: return {'toplam':len(data),'gecerli':len(valid),'hatali':len(errors),'hatalar':errors[:100]}
    if errors: raise HTTPException(422,{'mesaj':'Hatalı satırlar var; aktarım yapılmadı','hatalar':errors[:100]})
    count=0
    for line,r in valid:
        if entity=='istasyonlar':
            db.execute(text("INSERT INTO aftys.istasyon_birimler(tip,ad,kod,telsiz_adi) VALUES('ISTASYON',:a,:k,:ta)"),{'a':str(r['istasyon_adi']).strip(),'k':str(r.get('istasyon_telsiz_kodu') or '').strip() or None,'ta':str(r.get('istasyon_telsiz_adi') or '').strip() or None})
        elif entity=='servisler':
            db.execute(text("INSERT INTO aftys.servisler(servis_turu,ad,hizmet_verilen_markalar,eposta,adres) VALUES(:t,:a,:m,:e,:ad)"),{'t':r.get('_servis_turu') or 'ARAC','a':str(r['ad']).strip(),'m':r.get('hizmet_verilen_markalar'),'e':r.get('eposta'),'ad':r.get('adres')})
        elif entity=='kullanicilar':
            uid=uuid.uuid4(); tc=str(r['tc_kimlik_no']).strip(); adsoy=f"{str(r['ad']).strip()} {str(r['soyad']).strip()}"
            db.execute(text("INSERT INTO aftys.kullanicilar(kullanici_id,kullanici_adi,tc_kimlik_no,parola_hash,parola_degistir_zorunlu,ad_soyad,eposta,telefon,aktif_mi) VALUES(:id,:k,:tc,:p,true,:a,:e,:t,true)"),{'id':str(uid),'k':tc,'tc':tc,'p':hash_password('Degistir123!'),'a':adsoy,'e':r.get('eposta'),'t':r.get('telefon')})
        elif entity=='cagri_kayit':
            station=db.execute(text("SELECT birim_id FROM aftys.istasyon_birimler WHERE kod=:k AND aktif_mi=true"),{'k':str(r['istasyon_telsiz_kodu']).strip()}).scalar_one_or_none()
            if not station: raise HTTPException(422,f"İstasyon bulunamadı: {r['istasyon_telsiz_kodu']}")
            aid=None
            if r.get('plaka'):
                aid=db.execute(text("SELECT arac_id FROM aftys.arac_kimlik_gecmisi WHERE alan_tipi='PLAKA' AND upper(deger)=upper(:p) AND bitis_tarihi IS NULL"),{'p':str(r['plaka']).strip()}).scalar_one_or_none()
                if not aid: raise HTTPException(422,f"Plaka bulunamadı: {r['plaka']}")
            rawdt=r.get('cagri_tarihi_saati'); dt=rawdt if isinstance(rawdt,datetime) else datetime.strptime(str(rawdt).strip(),'%d.%m.%Y %H:%M')
            tt=str(r.get('talep_turu') or 'ARIZA').strip().upper(); status=str(r.get('durum') or ('GAYRIFAAL' if tt=='ARIZA' else 'YENI')).strip().upper(); no=next_number(db,'TLP',dt.year)
            sql="""INSERT INTO aftys.teknik_talepler(talep_no,istasyon_birim_id,bildiren_kisi,iletisim_yontemi,talep_turu,arac_id,aciklama,ariza_kategori,aciliyet,durum,created_by,created_at,updated_at,ariza_cozum_yeri,gayrifaal_baslangic,gayrifaal_bitis,kapanis_tarihi,kaynak) VALUES(:n,:b,:bk,'DIGER',:tt,:a,:ac,:kat,'NORMAL',:d,:u,:dt,:dt,:cy,CASE WHEN :fault THEN :dt ELSE NULL END,CASE WHEN :done THEN :dt ELSE NULL END,CASE WHEN :done THEN :dt ELSE NULL END,'TOPLU_AKTARIM') RETURNING talep_id"""
            tid=db.execute(text(sql),{'n':no,'b':station,'bk':r.get('bildiren_kisi') or None,'tt':tt,'a':str(aid) if aid else None,'ac':str(r['aciklama']).strip(),'kat':str(r.get('ariza_turu') or '').strip().upper() or None,'d':status,'u':str(u.kullanici_id),'dt':dt,'cy':str(r.get('ariza_cozum_yeri') or '').strip().upper() or None,'fault':tt=='ARIZA','done':status in ('TAMAMLANDI','IPTAL')}).scalar_one()
            db.execute(text("INSERT INTO aftys.teknik_talep_zaman_cizelgesi(talep_id,durum,aciklama,kullanici_id,created_at) VALUES(:id,:d,'Geçmiş çağrı kaydı toplu veri aktarımı ile eklendi',:u,:dt)"),{'id':tid,'d':status,'u':str(u.kullanici_id),'dt':dt})
        elif entity=='araclar':
            aid=uuid.uuid4(); station=None
            durum=db.execute(text("SELECT arac_durum_id FROM aftys.arac_durumlari WHERE kod='AKTIF_YEDEK'" )).scalar(); km=int(r.get('baslangic_km') or 0)
            db.execute(text("INSERT INTO aftys.araclar(arac_id,arac_sinifi,marka,model,model_yili,yakit_turu,arac_durum_id,gunluk_birim_id,gosterge_km,kumulatif_km,created_by) VALUES(:id,:s,:m,:mo,:y,:yt,:d,:b,:km,:km,:u)"),{'id':str(aid),'s':str(r['arac_sinifi']).strip(),'m':str(r['marka']).strip(),'mo':r.get('model'),'y':int(r['model_yili']) if r.get('model_yili') else None,'yt':r.get('yakit_turu'),'d':durum,'b':station,'km':km,'u':str(u.kullanici_id)})
            for typ,col in [('PLAKA','plaka'),('SASI_NO','sasi_no'),('MOTOR_NO','motor_no')]:
                if r.get(col): db.execute(text("INSERT INTO aftys.arac_kimlik_gecmisi(arac_id,alan_tipi,deger,degisiklik_tipi,created_by) VALUES(:a,:t,:v,'ILK_KAYIT',:u)"),{'a':str(aid),'t':typ,'v':str(r[col]).strip(),'u':str(u.kullanici_id)})
            db.execute(text("INSERT INTO aftys.km_gecmisi(arac_id,gosterge_km,kumulatif_km,tibbi_donanim_markasi,adblue_var,kaynak_tipi) VALUES(:a,:km,:km,:tm,:ad,'ARAC_ILK_KAYIT')"),{'a':str(aid),'km':km,'tm':r.get('tibbi_donanim_markasi'),'ad':str(r.get('adblue_var') or '').upper() in ('EVET','TRUE','1','VAR')})
        count+=1
    logid=db.execute(text("INSERT INTO aftys.toplu_aktarimlar(kullanici_id,varlik_turu,dosya_adi,kayit_sayisi,durum) VALUES(:u,:e,:d,:c,'TAMAMLANDI') RETURNING aktarim_id"),{'u':str(u.kullanici_id),'e':entity,'d':dosya.filename or 'dosya.xlsx','c':count}).scalar_one(); audit(db,u,'TOPLU_AKTARIM',logid,'IMPORT',{'entity':entity,'count':count}); db.commit(); return {'aktarim_id':logid,'kayit_sayisi':count}

@router.get('/files/{ek_id}')
def get_file(ek_id:uuid.UUID,db:Session=Depends(get_db),u=Depends(current_user)):
    r=db.execute(text("SELECT dosya_adi,mime_type,storage_key FROM aftys.dosyalar_ekler WHERE ek_id=:id"),{'id':str(ek_id)}).mappings().first()
    if not r: raise HTTPException(404,'Dosya bulunamadı')
    p=pathlib.Path(os.environ.get('AFTYS_UPLOAD_DIR','/data/uploads'))/r['storage_key']
    if not p.exists(): raise HTTPException(404,'Dosya depoda bulunamadı')
    return FileResponse(str(p),media_type=r['mime_type'] or 'application/octet-stream',filename=r['dosya_adi'])

@router.get('/talepler/{talep_id}/detay')
def request_detail(talep_id:int,db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Teknik Destek Talepleri','GORUNTULE')
    t=db.execute(text("""SELECT t.*,trim(concat(COALESCE(ib.kod,''),CASE WHEN ib.kod IS NOT NULL AND ib.telsiz_adi IS NOT NULL THEN ' ' ELSE '' END,COALESCE(ib.telsiz_adi,''),' — ',ib.ad)) istasyon_birim,
      (SELECT deger FROM aftys.arac_kimlik_gecmisi WHERE arac_id=t.arac_id AND alan_tipi='PLAKA' AND bitis_tarihi IS NULL LIMIT 1) plaka
      FROM aftys.teknik_talepler t LEFT JOIN aftys.istasyon_birimler ib ON ib.birim_id=t.istasyon_birim_id WHERE t.talep_id=:id"""),{'id':talep_id}).mappings().first()
    if not t: raise HTTPException(404,'Talep bulunamadı')
    timeline=rows(db.execute(text("""SELECT z.*,k.ad_soyad FROM aftys.teknik_talep_zaman_cizelgesi z LEFT JOIN aftys.kullanicilar k ON k.kullanici_id=z.kullanici_id WHERE z.talep_id=:id ORDER BY z.created_at"""),{'id':talep_id}))
    return {'talep':dict(t),'timeline':timeline}

@router.post('/talepler/{talep_id}/durum')
def request_status(talep_id:int,b:dict[str,Any]=Body(...),db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Teknik Destek Talepleri','DUZENLE')
    allowed={'GAYRIFAAL','SERVISE_GIRIS','ARIZA_TESPIT','PARCA_BEKLIYOR','ONARIM','SERVISTEN_CIKIS','TEST_KONTROL','TAMAMLANDI','IPTAL'}
    d=str(b.get('durum') or ''); note=str(b.get('aciklama') or '').strip()
    if d not in allowed: raise HTTPException(422,'Geçersiz durum')
    if d in {'PARCA_BEKLIYOR','IPTAL'} and not note: raise HTTPException(422,'Bu durum için açıklama zorunlu')
    db.execute(text("UPDATE aftys.teknik_talepler SET durum=:d,durum_notu=:n,updated_at=now(),kapanis_tarihi=CASE WHEN CAST(:d_status AS text) IN ('TAMAMLANDI','IPTAL') THEN now() ELSE NULL END,gayrifaal_bitis=CASE WHEN CAST(:d_status AS text)='TAMAMLANDI' THEN now() ELSE gayrifaal_bitis END WHERE talep_id=:id"),{'d':d,'d_status':d,'n':note or None,'id':talep_id})
    if d=='TAMAMLANDI':
        aid=db.execute(text("SELECT arac_id FROM aftys.teknik_talepler WHERE talep_id=:id"),{'id':talep_id}).scalar_one_or_none()
        if aid:
            code='AKTIF_GOREVDE' if db.execute(text("SELECT gunluk_birim_id FROM aftys.araclar WHERE arac_id=:a"),{'a':str(aid)}).scalar_one_or_none() else 'AKTIF_YEDEK'
            did=db.execute(text("SELECT arac_durum_id FROM aftys.arac_durumlari WHERE kod=:k"),{'k':code}).scalar_one_or_none()
            if did: db.execute(text("UPDATE aftys.araclar SET arac_durum_id=:d WHERE arac_id=:a"),{'d':did,'a':str(aid)})
    db.execute(text("INSERT INTO aftys.teknik_talep_zaman_cizelgesi(talep_id,durum,aciklama,kullanici_id) VALUES(:id,:d,:n,:u)"),{'id':talep_id,'d':d,'n':note or None,'u':str(u.kullanici_id)})
    audit(db,u,'TEKNIK_TALEP',talep_id,'DURUM_DEGISTIR',{'durum':d,'aciklama':note}); db.commit(); return {'ok':True}

@router.post('/talepler/{talep_id}/degisiklik-talebi',status_code=201)
def request_change_approval(talep_id:int,b:dict[str,Any]=Body(...),db:Session=Depends(get_db),u=Depends(current_user)):
    tur=str(b.get('islem_turu') or 'DUZENLE'); gerekce=str(b.get('gerekce') or '').strip()
    if tur not in ('DUZENLE','IPTAL'): raise HTTPException(422,'İşlem türü geçersiz')
    if not gerekce: raise HTTPException(422,'Gerekçe zorunlu')
    oid=db.execute(text("INSERT INTO aftys.onay_talepleri(entity_type,entity_id,islem_turu,yeni_json,gerekce,talep_eden) VALUES('TEKNIK_TALEP',:id,:t,CAST(:j AS jsonb),:g,:u) RETURNING onay_id"),{'id':str(talep_id),'t':tur,'j':json.dumps(b.get('degisiklikler') or {},ensure_ascii=False),'g':gerekce,'u':str(u.kullanici_id)}).scalar_one()
    db.execute(text("""INSERT INTO aftys.bildirimler(kullanici_id,baslik,mesaj,link) SELECT DISTINCT kr.kullanici_id,'Onay Bekleyen İşlem',:m,:l FROM aftys.kullanici_rolleri kr JOIN aftys.roller r ON r.rol_id=kr.rol_id WHERE r.ad IN ('Sistem Yöneticisi','Yönetici')"""),{'m':f'Teknik destek kaydı #{talep_id} için {tur} talebi oluşturuldu.','l':f'/onay/{oid}'})
    db.commit(); return {'onay_id':oid,'durum':'BEKLIYOR'}

@router.delete('/talepler/{talep_id}')
def delete_request(talep_id:int,b:dict[str,Any]=Body(...),db:Session=Depends(get_db),u=Depends(current_user)):
    gerekce=str(b.get('gerekce') or '').strip()
    if not gerekce: raise HTTPException(422,'Silme gerekçesi zorunlu')
    row=db.execute(text("SELECT talep_no FROM aftys.teknik_talepler WHERE talep_id=:id AND COALESCE(pasif_mi,false)=false"),{'id':talep_id}).first()
    if not row: raise HTTPException(404,'Kayıt bulunamadı')
    roles=[x[0] for x in db.execute(text("SELECT r.ad FROM aftys.kullanici_rolleri kr JOIN aftys.roller r ON r.rol_id=kr.rol_id WHERE kr.kullanici_id=:u"),{'u':str(u.kullanici_id)}).all()]
    if any(x in ('Sistem Yöneticisi','Yönetici') for x in roles):
        db.execute(text("UPDATE aftys.teknik_talepler SET pasif_mi=true,durum='IPTAL',kapanis_tarihi=now(),updated_at=now() WHERE talep_id=:id"),{'id':talep_id})
        db.execute(text("INSERT INTO aftys.teknik_talep_zaman_cizelgesi(talep_id,durum,aciklama,kullanici_id) VALUES(:id,'IPTAL',:n,:u)"),{'id':talep_id,'n':'Yönetici tarafından silindi/pasife alındı: '+gerekce,'u':str(u.kullanici_id)})
        audit(db,u,'TEKNIK_TALEP',talep_id,'DELETE_SOFT',gerekce=gerekce); db.commit(); return {'ok':True,'sonuc':'SILINDI'}
    oid=db.execute(text("INSERT INTO aftys.onay_talepleri(entity_type,entity_id,islem_turu,gerekce,talep_eden) VALUES('TEKNIK_TALEP',:id,'SIL',:g,:u) RETURNING onay_id"),{'id':str(talep_id),'g':gerekce,'u':str(u.kullanici_id)}).scalar_one()
    db.commit(); return {'ok':True,'sonuc':'ONAYA_GONDERILDI','onay_id':oid}

@router.get('/bildirimler')
def notifications(db:Session=Depends(get_db),u=Depends(current_user)):
    return {'okunmamis':db.execute(text("SELECT count(*) FROM aftys.bildirimler WHERE kullanici_id=:u AND okundu=false"),{'u':str(u.kullanici_id)}).scalar(),'records':rows(db.execute(text("SELECT * FROM aftys.bildirimler WHERE kullanici_id=:u ORDER BY created_at DESC LIMIT 100"),{'u':str(u.kullanici_id)}))}

# ======================= V10 =======================
@router.get('/zimmetler')
def assignments(db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Araçlar','GORUNTULE')
    return rows(db.execute(text("""SELECT t.tahsis_id,t.arac_id::text,t.birim_id,t.baslangic_tarihi,t.bitis_tarihi,t.neden,t.aciklama,
      COALESCE((SELECT deger FROM aftys.arac_kimlik_gecmisi WHERE arac_id=a.arac_id AND alan_tipi='PLAKA' AND bitis_tarihi IS NULL LIMIT 1),'—') plaka,
      a.marka,a.model,ib.ad istasyon_adi,ib.kod telsiz_kodu,ib.telsiz_adi
      FROM aftys.arac_tahsisleri t JOIN aftys.araclar a ON a.arac_id=t.arac_id JOIN aftys.istasyon_birimler ib ON ib.birim_id=t.birim_id
      WHERE t.bitis_tarihi IS NULL ORDER BY COALESCE(ib.kod,''),ib.ad""")))

@router.post('/zimmetler',status_code=201)
def assign_vehicle(b:dict[str,Any]=Body(...),db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Araçlar','DUZENLE')
    aid=str(b.get('arac_id') or ''); bid=b.get('birim_id')
    if not aid or not bid: raise HTTPException(422,'Araç ve istasyon/birim zorunlu')
    db.execute(text("SELECT pg_advisory_xact_lock(:k)"),{'k':int(bid)})
    if not db.execute(text("SELECT 1 FROM aftys.araclar WHERE arac_id=:a AND aktif_mi=true AND arac_sinifi='AMBULANS'"),{'a':aid}).first(): raise HTTPException(404,'Aktif ambulans bulunamadı')
    occupied=db.execute(text("SELECT t.tahsis_id,(SELECT deger FROM aftys.arac_kimlik_gecmisi WHERE arac_id=t.arac_id AND alan_tipi='PLAKA' AND bitis_tarihi IS NULL LIMIT 1) plaka FROM aftys.arac_tahsisleri t WHERE t.birim_id=:b AND t.bitis_tarihi IS NULL AND t.arac_id<>CAST(:a AS uuid) LIMIT 1"),{'b':bid,'a':aid}).mappings().first()
    if occupied: raise HTTPException(409,f"Bu istasyonda zaten aktif zimmetli ambulans var: {occupied['plaka'] or 'PLAKASIZ'}")
    db.execute(text("UPDATE aftys.arac_tahsisleri SET bitis_tarihi=GREATEST(now(),baslangic_tarihi + interval '1 second'),updated_at=now() WHERE arac_id=:a AND bitis_tarihi IS NULL"),{'a':aid})
    tid=db.execute(text("INSERT INTO aftys.arac_tahsisleri(arac_id,birim_id,baslangic_tarihi,neden,aciklama,created_by) VALUES(:a,:b,COALESCE(CAST(:t AS timestamptz),now()),:n,:ac,:u) RETURNING tahsis_id"),{'a':aid,'b':bid,'t':b.get('baslangic_tarihi') or None,'n':b.get('neden') or 'İstasyon zimmeti','ac':b.get('aciklama') or None,'u':str(u.kullanici_id)}).scalar_one()
    db.execute(text("UPDATE aftys.araclar SET gunluk_birim_id=:b WHERE arac_id=:a"),{'b':bid,'a':aid})
    audit(db,u,'ARAC_ZIMMET',tid,'CREATE',{'arac_id':aid,'birim_id':bid}); db.commit(); return {'tahsis_id':tid}

@router.post('/zimmetler/{tahsis_id}/kaldir')
def unassign_vehicle(tahsis_id:int,b:dict[str,Any]=Body(default={}),db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Araçlar','DUZENLE')
    r=db.execute(text("SELECT arac_id FROM aftys.arac_tahsisleri WHERE tahsis_id=:id AND bitis_tarihi IS NULL"),{'id':tahsis_id}).first()
    if not r: raise HTTPException(404,'Aktif zimmet bulunamadı')
    db.execute(text("UPDATE aftys.arac_tahsisleri SET bitis_tarihi=GREATEST(now(),baslangic_tarihi + interval '1 second'),bitis_km=:km,aciklama=concat_ws(E'\\n',aciklama,CAST(:ac AS TEXT)),updated_at=now() WHERE tahsis_id=:id"),{'id':tahsis_id,'km':b.get('bitis_km') or None,'ac':b.get('aciklama') or 'Zimmet kaldırıldı'})
    db.execute(text("UPDATE aftys.araclar SET gunluk_birim_id=NULL WHERE arac_id=:a"),{'a':str(r[0])}); audit(db,u,'ARAC_ZIMMET',tahsis_id,'CLOSE',b); db.commit(); return {'ok':True}

@router.get('/zimmetler/by-station/{birim_id}')
def assignment_by_station(birim_id:int,db:Session=Depends(get_db),u=Depends(current_user)):
    require_permission(db,u,'Araçlar','GORUNTULE')
    r=db.execute(text("""SELECT t.arac_id::text arac_id,(SELECT deger FROM aftys.arac_kimlik_gecmisi WHERE arac_id=t.arac_id AND alan_tipi='PLAKA' AND bitis_tarihi IS NULL LIMIT 1) plaka FROM aftys.arac_tahsisleri t WHERE t.birim_id=:b AND t.bitis_tarihi IS NULL ORDER BY t.baslangic_tarihi DESC LIMIT 1"""),{'b':birim_id}).mappings().first()
    return dict(r) if r else None

@router.get('/onay-talepleri')
def approvals(db:Session=Depends(get_db),u=Depends(current_user)):
    roles=[x[0] for x in db.execute(text("SELECT r.ad FROM aftys.kullanici_rolleri kr JOIN aftys.roller r ON r.rol_id=kr.rol_id WHERE kr.kullanici_id=:u"),{'u':str(u.kullanici_id)}).all()]
    if not any(x in ('Sistem Yöneticisi','Yönetici') for x in roles): raise HTTPException(403,'Yönetici yetkisi gerekir')
    return rows(db.execute(text("""SELECT o.*,k.ad_soyad talep_eden_adi FROM aftys.onay_talepleri o LEFT JOIN aftys.kullanicilar k ON k.kullanici_id=o.talep_eden WHERE o.durum='BEKLIYOR' ORDER BY o.created_at DESC""")))

@router.post('/onay-talepleri/{onay_id}/karar')
def approval_decision(onay_id:int,b:dict[str,Any]=Body(...),db:Session=Depends(get_db),u=Depends(current_user)):
    roles=[x[0] for x in db.execute(text("SELECT r.ad FROM aftys.kullanici_rolleri kr JOIN aftys.roller r ON r.rol_id=kr.rol_id WHERE kr.kullanici_id=:u"),{'u':str(u.kullanici_id)}).all()]
    if not any(x in ('Sistem Yöneticisi','Yönetici') for x in roles): raise HTTPException(403,'Yönetici yetkisi gerekir')
    karar=str(b.get('karar') or '').upper(); row=db.execute(text("SELECT * FROM aftys.onay_talepleri WHERE onay_id=:id AND durum='BEKLIYOR'"),{'id':onay_id}).mappings().first()
    if not row: raise HTTPException(404,'Bekleyen onay bulunamadı')
    if karar not in ('ONAYLANDI','REDDEDILDI'): raise HTTPException(422,'Karar geçersiz')
    if karar=='ONAYLANDI' and row['entity_type']=='TEKNIK_TALEP' and row['islem_turu'] in ('IPTAL','SIL'):
        db.execute(text("UPDATE aftys.teknik_talepler SET pasif_mi=true,durum='IPTAL',kapanis_tarihi=now(),updated_at=now() WHERE talep_id=:id"),{'id':int(row['entity_id'])})
        db.execute(text("INSERT INTO aftys.teknik_talep_zaman_cizelgesi(talep_id,durum,aciklama,kullanici_id) VALUES(:id,'IPTAL',:n,:u)"),{'id':int(row['entity_id']),'n':'Admin onayı ile pasife alındı. '+str(row['gerekce'] or ''),'u':str(u.kullanici_id)})
    db.execute(text("UPDATE aftys.onay_talepleri SET durum=:d,karar_veren=:u,karar_notu=:n,karar_tarihi=now() WHERE onay_id=:id"),{'d':karar,'u':str(u.kullanici_id),'n':b.get('not') or None,'id':onay_id}); audit(db,u,'ONAY_TALEBI',onay_id,'DECISION',{'karar':karar}); db.commit(); return {'ok':True}
